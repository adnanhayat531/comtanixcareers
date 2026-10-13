"""Extract contact fields from an uploaded CV (PDF / DOCX)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

from fastapi import HTTPException, UploadFile, status

EMAIL_RE = re.compile(r"\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}\b")
PHONE_RE = re.compile(
    r"(?:\+?92[\s\-]?)?(?:0)?3\d{2}[\s\-]?\d{7}"
    r"|(?:\+?\d{1,3}[\s\-]?)?(?:\(?\d{2,4}\)?[\s\-]?)?\d{3,4}[\s\-]?\d{4,7}"
)
NAME_SKIP = re.compile(
    r"^(curriculum|vitae|resume|cv|profile|contact|email|phone|mobile|"
    r"address|objective|summary|experience|education|skills)\b",
    re.I,
)


@dataclass
class ParsedCvFields:
    name: str = ""
    email: str = ""
    phone: str = ""
    role: str = ""
    exp: str = ""


def _extract_pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="PDF parsing is not available on the server.",
        ) from exc

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages[:6]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _extract_docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Word parsing is not available on the server.",
        ) from exc

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs[:120] if p.text]
    for table in doc.tables[:8]:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def _guess_name(text: str) -> str:
    for raw in text.splitlines()[:12]:
        line = " ".join(raw.split()).strip(" -|•\t")
        if not line or len(line) < 3 or len(line) > 60:
            continue
        if EMAIL_RE.search(line) or PHONE_RE.search(line):
            continue
        if NAME_SKIP.search(line):
            continue
        if any(ch.isdigit() for ch in line):
            continue
        words = line.split()
        if 2 <= len(words) <= 5 and all(w[0].isalpha() for w in words if w):
            return line.title() if line.isupper() else line
        if len(words) == 1 and line.isalpha() and line[:1].isupper():
            return line
    return ""


def _guess_role(text: str) -> str:
    for raw in text.splitlines()[:20]:
        line = " ".join(raw.split()).strip()
        if not line or len(line) > 80:
            continue
        lower = line.lower()
        if any(
            key in lower
            for key in (
                "engineer",
                "developer",
                "analyst",
                "support",
                "administrator",
                "manager",
                "specialist",
                "associate",
                "intern",
            )
        ):
            if not EMAIL_RE.search(line) and not NAME_SKIP.search(line):
                return line
    return "General / future openings"


def _guess_exp(text: str) -> str:
    patterns = [
        r"(\d{1,2})\+?\s*(?:years?|yrs?)\s+(?:of\s+)?experience",
        r"experience\s*[:\-]?\s*(\d{1,2})\+?\s*(?:years?|yrs?)",
        r"\b(fresher|fresh graduate|entry[- ]level|internship)\b",
    ]
    lower = text.lower()
    for pat in patterns:
        m = re.search(pat, lower, re.I)
        if not m:
            continue
        if m.lastindex and m.group(1) and m.group(1).isdigit():
            years = int(m.group(1))
            if years < 1:
                return "Less than 1 year"
            if years <= 3:
                return "1 to 3 years"
            return "3 or more years"
        token = m.group(0).lower()
        if "fresh" in token or "entry" in token or "intern" in token:
            return "Fresher"
    return ""


def parse_cv_bytes(filename: str, data: bytes) -> ParsedCvFields:
    name_lower = (filename or "").lower()
    if name_lower.endswith(".pdf"):
        text = _extract_pdf_text(data)
    elif name_lower.endswith(".docx"):
        text = _extract_docx_text(data)
    elif name_lower.endswith(".doc"):
        # Legacy .doc is binary; best-effort latin decode for embedded strings.
        text = data.decode("latin-1", errors="ignore")
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Please upload a PDF or Word file (.pdf, .doc, .docx).",
        )

    text = text.replace("\x00", " ")
    email_match = EMAIL_RE.search(text)
    phone_match = PHONE_RE.search(text)
    phone = ""
    if phone_match:
        phone = re.sub(r"[^\d+]", "", phone_match.group(0))
        if phone.startswith("03") and len(phone) == 11:
            phone = phone
        elif phone.startswith("923") and len(phone) >= 12:
            phone = "0" + phone[2:]

    return ParsedCvFields(
        name=_guess_name(text),
        email=(email_match.group(0) if email_match else "").lower(),
        phone=phone,
        role=_guess_role(text),
        exp=_guess_exp(text),
    )


async def parse_uploaded_cv(upload: UploadFile) -> ParsedCvFields:
    data = await upload.read()
    if not data:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Please upload your CV to continue.",
        )
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Your CV must be 5 MB or smaller.",
        )
    return parse_cv_bytes(upload.filename or "resume.pdf", data)
